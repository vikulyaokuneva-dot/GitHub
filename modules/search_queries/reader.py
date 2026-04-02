"""Reader for WB search queries XLSX report."""

from __future__ import annotations

from pathlib import Path
from typing import Any


SHEET_CANDIDATES: tuple[str, ...] = (
    "Детальные показатели",
    # Safe fallback for already existing WB exports in repo.
    "Статистика по ключевым словам",
)

HEADER_ALIASES: dict[str, tuple[str, ...]] = {
    "seller_sku": ("Артикул продавца", "Артикул продавца, штрихкод"),
    "wb_sku": ("Артикул WB", "Номенклатура WB", "nmId"),
    "query": ("Поисковый запрос", "Ключевая фраза"),
    "query_count": ("Количество запросов",),
    "visibility_pct": ("Видимость, %", "Видимость,%"),
    "avg_position": ("Средняя позиция",),
    "median_position": ("Медианная позиция",),
    "card_clicks": ("Переходы в карточку",),
    "add_to_cart": ("Положили в корзину",),
    "cart_conv_pct": ("Конверсия в корзину, %", "Конверсия в корзину,%"),
    "orders_count": ("Заказали, шт",),
    "order_conv_pct": ("Конверсия в заказ, %", "Конверсия в заказ,%"),
    "price_min": ("Минимальная цена со скидкой",),
    "price_max": ("Максимальная цена со скидкой",),
}


def _pick_sheet_name(sheet_names: list[str]) -> str | None:
    for candidate in SHEET_CANDIDATES:
        if candidate in sheet_names:
            return candidate
    return None


def _score_header(values: list[str]) -> int:
    values_set = {x.strip().lower() for x in values if x and x.strip()}
    score = 0
    for aliases in HEADER_ALIASES.values():
        for alias in aliases:
            if alias.strip().lower() in values_set:
                score += 1
                break
    return score


def _find_header_row(rows: list[list[Any]]) -> int | None:
    best_idx = None
    best_score = 0
    for idx, row in enumerate(rows):
        stringified = [str(v).strip() if v is not None else "" for v in row]
        score = _score_header(stringified)
        if score > best_score:
            best_score = score
            best_idx = idx
    if best_idx is None:
        return None
    return best_idx if best_score > 0 else None


def read_search_queries_xlsx(path: str | Path) -> dict[str, Any]:
    """
    Read WB search queries XLSX and return raw row dictionaries.

    Returns dict:
      status: ok | no_input | dependency_missing | read_error | sheet_not_found | header_not_found
      rows: list[dict]
      sheet_name: str | None
      message: str
    """
    xlsx_path = Path(path)
    if not xlsx_path.exists():
        return {
            "status": "no_input",
            "rows": [],
            "sheet_name": None,
            "message": f"Input file not found: {xlsx_path}",
        }

    try:
        from openpyxl import load_workbook  # type: ignore
    except Exception as exc:  # pragma: no cover
        return {
            "status": "dependency_missing",
            "rows": [],
            "sheet_name": None,
            "message": f"openpyxl is required to read xlsx ({exc})",
        }

    try:
        wb = load_workbook(xlsx_path, read_only=True, data_only=True)
    except Exception as exc:
        return {
            "status": "read_error",
            "rows": [],
            "sheet_name": None,
            "message": f"Failed to read xlsx: {exc}",
        }

    sheet_name = _pick_sheet_name(list(wb.sheetnames))
    if sheet_name is None:
        return {
            "status": "sheet_not_found",
            "rows": [],
            "sheet_name": None,
            "message": (
                "Sheet not found. Expected one of: "
                + ", ".join(SHEET_CANDIDATES)
                + f". Available: {', '.join(wb.sheetnames)}"
            ),
        }

    ws = wb[sheet_name]
    all_rows: list[list[Any]] = [list(r) for r in ws.iter_rows(values_only=True)]
    if not all_rows:
        return {
            "status": "ok",
            "rows": [],
            "sheet_name": sheet_name,
            "message": "Sheet is empty",
        }

    scan_rows = all_rows[:50]
    header_idx = _find_header_row(scan_rows)
    if header_idx is None:
        return {
            "status": "header_not_found",
            "rows": [],
            "sheet_name": sheet_name,
            "message": "Header row not found in first 50 rows",
        }

    headers = [str(v).strip() if v is not None else "" for v in all_rows[header_idx]]
    data_rows = all_rows[header_idx + 1 :]
    raw_rows: list[dict[str, Any]] = []
    for row in data_rows:
        if not any(cell is not None and str(cell).strip() != "" for cell in row):
            continue
        item: dict[str, Any] = {}
        max_cols = max(len(headers), len(row))
        for idx in range(max_cols):
            key = headers[idx] if idx < len(headers) else f"col_{idx + 1}"
            if not key:
                key = f"col_{idx + 1}"
            val = row[idx] if idx < len(row) else None
            item[key] = val
        raw_rows.append(item)

    return {
        "status": "ok",
        "rows": raw_rows,
        "sheet_name": sheet_name,
        "message": f"Loaded {len(raw_rows)} rows from {sheet_name}",
    }

