"""COGS loader for seller-scoped XLSX files."""

from __future__ import annotations

import re
from pathlib import Path
from typing import Any


COGS_HEADER_ALIASES: dict[str, tuple[str, ...]] = {
    "seller_sku": (
        "seller_sku",
        "sku продавца",
        "артикул продавца",
        "артикул поставщика",
    ),
    "cogs": (
        "cogs",
        "себестоимость",
        "себестоимость за 1 шт",
        "cost",
        "cost_price",
    ),
}


def _norm(text: Any) -> str:
    s = "" if text is None else str(text)
    s = s.replace("\xa0", " ").replace("\n", " ").replace("\r", " ")
    s = s.strip().lower().replace("ё", "е")
    s = re.sub(r"\s+", " ", s)
    return s


def _to_float(value: Any) -> float | None:
    if value is None:
        return None
    if isinstance(value, (int, float)):
        return float(value)
    s = str(value).strip().replace(" ", "").replace(",", ".")
    if not s:
        return None
    try:
        return float(s)
    except Exception:
        return None


def _match_field(header: Any) -> str | None:
    h = _norm(header)
    if not h:
        return None
    best: str | None = None
    best_score = -1
    for field, aliases in COGS_HEADER_ALIASES.items():
        for alias in aliases:
            a = _norm(alias)
            if not a:
                continue
            if h == a or a in h:
                score = len(a) + (100 if h == a else 0)
                if score > best_score:
                    best = field
                    best_score = score
    return best


def load_cogs_xlsx(path: str | Path) -> dict[str, Any]:
    """Load COGS map from runtime/cabinets/<seller>/raw/cogs.xlsx."""
    source = Path(path)
    warnings: list[str] = []
    out: dict[str, float] = {}

    if not source.exists():
        return {
            "status": "no_input",
            "message": f"COGS file not found: {source}",
            "items": out,
            "warnings": warnings,
        }

    try:
        from openpyxl import load_workbook  # type: ignore
    except Exception as exc:  # pragma: no cover
        return {
            "status": "dependency_missing",
            "message": f"openpyxl is required to read cogs.xlsx ({exc})",
            "items": out,
            "warnings": warnings,
        }

    try:
        wb = load_workbook(source, read_only=False, data_only=True)
    except Exception as exc:
        return {
            "status": "read_error",
            "message": f"Failed to read COGS xlsx: {exc}",
            "items": out,
            "warnings": warnings,
        }

    ws = wb[wb.sheetnames[0]]
    rows = [list(r) for r in ws.iter_rows(values_only=True)]
    if not rows:
        return {
            "status": "ok",
            "message": "COGS sheet is empty",
            "items": out,
            "warnings": warnings,
        }

    header_idx = None
    header_map: dict[str, int] = {}
    max_scan = min(len(rows), 150)
    for idx in range(max_scan):
        row = rows[idx]
        detected: dict[str, int] = {}
        for col_i, cell in enumerate(row):
            field = _match_field(cell)
            if field and field not in detected:
                detected[field] = col_i
        if "seller_sku" in detected and "cogs" in detected:
            header_idx = idx
            header_map = detected
            break

    if header_idx is None:
        return {
            "status": "header_not_found",
            "message": "COGS header not found (required: seller_sku, cogs)",
            "items": out,
            "warnings": warnings,
        }

    bad_rows = 0
    for row in rows[header_idx + 1 :]:
        sku_col = header_map["seller_sku"]
        cogs_col = header_map["cogs"]
        sku = _norm(row[sku_col] if sku_col < len(row) else None)
        cogs_raw = row[cogs_col] if cogs_col < len(row) else None
        cogs = _to_float(cogs_raw)
        if not sku:
            continue
        if cogs is None:
            bad_rows += 1
            continue
        out[sku] = float(cogs)

    if bad_rows > 0:
        warnings.append(f"Skipped {bad_rows} COGS rows with invalid cogs values")

    return {
        "status": "ok",
        "message": f"Loaded {len(out)} COGS rows from {source.name}",
        "items": out,
        "warnings": warnings,
    }

