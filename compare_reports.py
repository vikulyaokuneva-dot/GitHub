#!/usr/bin/env python3
"""Сравнить структуры отчетов v2 vs v5"""

import json
import zipfile
from pathlib import Path
from pprint import pprint

project_dir = Path(r"D:\WB\Бот ИИ менеджер\GitHub")

# Проверим v5 JSON файлы
print("=" * 60)
print("V5 CURRENT METRICS")
print("=" * 60)

metrics_v5_file = project_dir / "cabinets" / "seller_001" / "artifacts" / "metrics.json"
if metrics_v5_file.exists():
    with open(metrics_v5_file) as f:
        metrics_v5 = json.load(f)
    print(f"Total metrics fields: {len(metrics_v5)}")
    print(f"Main keys: {list(metrics_v5.keys())[:10]}")
    
    # Посмотрим структуру
    if 'sku_metrics' in metrics_v5:
        print(f"SKU metrics count: {len(metrics_v5['sku_metrics'])}")
    if 'totals' in metrics_v5:
        print(f"Totals: {metrics_v5['totals']}")
    if 'financial_kpi' in metrics_v5:
        print(f"Financial KPI: {metrics_v5['financial_kpi']}")

print("\n" + "=" * 60)
print("V2 REPORT STRUCTURE")
print("=" * 60)

# Посмотрим v2 отчет
v2_zip = project_dir / "v2-daily-out (3).zip"
if v2_zip.exists():
    with zipfile.ZipFile(v2_zip) as z:
        print(f"Files in v2: {z.namelist()}")
        
        if 'metrics.json' in z.namelist():
            metrics_v2 = json.loads(z.read('metrics.json'))
            print(f"\nV2 Metrics keys: {list(metrics_v2.keys())[:20]}")
            
            # Посмотрим структуру данных
            if 'sku_metrics' in metrics_v2:
                print(f"V2 SKU metrics count: {len(metrics_v2['sku_metrics'])}")
                if metrics_v2['sku_metrics']:
                    first_sku = metrics_v2['sku_metrics'][0]
                    print(f"V2 SKU structure: {list(first_sku.keys())}")
                    
            if 'totals' in metrics_v2:
                print(f"V2 Totals: {metrics_v2['totals']}")

# Посмотрим реальные отчеты WB
print("\n" + "=" * 60)
print("REAL WB REPORTS")
print("=" * 60)

import openpyxl

excel_files = list(project_dir.glob("*.xlsx"))
print(f"Excel files found: {len(excel_files)}")

for xlsx in excel_files:
    print(f"\n  {xlsx.name}")
    try:
        wb = openpyxl.load_workbook(xlsx)
        print(f"    Sheets: {wb.sheetnames}")
        for sheet_name in wb.sheetnames[:3]:
            ws = wb[sheet_name]
            print(f"    Sheet '{sheet_name}': {ws.max_row} rows x {ws.max_column} cols")
    except Exception as e:
        print(f"    Error: {e}")
