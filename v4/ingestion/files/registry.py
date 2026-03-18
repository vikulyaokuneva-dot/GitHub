"""File-ingestion registry and helpers for audit mode.

Input: local file paths.
Output: source classification and raw tabular rows.
Does not compute KPI.
"""

from __future__ import annotations

import csv
import json
from dataclasses import dataclass
from pathlib import Path
from typing import Any


@dataclass(frozen=True)
class FileSourceDescriptor:
    source_name: str
    required: bool
    filename_keywords: tuple[str, ...]
    extensions: tuple[str, ...]


FILE_SOURCE_REGISTRY: tuple[FileSourceDescriptor, ...] = (
    FileSourceDescriptor(
        source_name="daily_report",
        required=True,
        filename_keywords=("daily", "supplier", "realization", "sales", "днев", "ежеднев", "детал"),
        extensions=(".json", ".csv", ".tsv", ".xlsx", ".xlsm"),
    ),
    FileSourceDescriptor(
        source_name="funnel_report",
        required=False,
        filename_keywords=("funnel", "ворон", "card", "conversion"),
        extensions=(".json", ".csv", ".tsv", ".xlsx", ".xlsm"),
    ),
    FileSourceDescriptor(
        source_name="ads_report",
        required=False,
        filename_keywords=("ads", "advert", "promo", "реклам"),
        extensions=(".json", ".csv", ".tsv", ".xlsx", ".xlsm"),
    ),
)


def get_required_file_inputs() -> tuple[str, ...]:
    return tuple(descriptor.source_name for descriptor in FILE_SOURCE_REGISTRY if descriptor.required)


def get_optional_file_inputs() -> tuple[str, ...]:
    return tuple(descriptor.source_name for descriptor in FILE_SOURCE_REGISTRY if not descriptor.required)


def is_allowed_source(source_name: str) -> bool:
    names = {descriptor.source_name for descriptor in FILE_SOURCE_REGISTRY}
    return str(source_name).strip() in names


def supported_extensions() -> tuple[str, ...]:
    values: list[str] = []
    for descriptor in FILE_SOURCE_REGISTRY:
        values.extend(descriptor.extensions)
    seen: set[str] = set()
    out: list[str] = []
    for value in values:
        if value in seen:
            continue
        seen.add(value)
        out.append(value)
    return tuple(out)


def _name_matches(path: Path, descriptor: FileSourceDescriptor) -> bool:
    stem = path.stem.lower()
    if path.suffix.lower() not in descriptor.extensions:
        return False
    return any(keyword in stem for keyword in descriptor.filename_keywords)


def classify_input_files(paths: list[Path]) -> dict[str, Path]:
    selected: dict[str, Path] = {}
    normalized_paths = sorted({path.resolve() for path in paths if path.exists()}, key=lambda p: str(p))
    for descriptor in FILE_SOURCE_REGISTRY:
        for path in normalized_paths:
            if _name_matches(path, descriptor):
                selected[descriptor.source_name] = path
                break
    return selected


def _read_json_rows(path: Path) -> list[dict[str, Any]]:
    payload = json.loads(path.read_text(encoding="utf-8"))
    if isinstance(payload, list):
        return [row for row in payload if isinstance(row, dict)]
    if isinstance(payload, dict):
        rows = payload.get("rows")
        if isinstance(rows, list):
            return [row for row in rows if isinstance(row, dict)]
    return []


def _read_csv_rows(path: Path, delimiter: str) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    with path.open("r", encoding="utf-8-sig", newline="") as file:
        reader = csv.DictReader(file, delimiter=delimiter)
        for row in reader:
            rows.append(dict(row))
    return rows


def _normalize_header(value: object, idx: int) -> str:
    if value is None:
        return f"col_{idx + 1}"
    text = str(value).strip()
    if not text:
        return f"col_{idx + 1}"
    return text


def _read_xlsx_rows(path: Path) -> list[dict[str, Any]]:
    try:
        from openpyxl import load_workbook  # type: ignore
    except Exception:
        raise RuntimeError("openpyxl is required to parse xlsx files")

    workbook = load_workbook(filename=str(path), data_only=True, read_only=True)
    try:
        sheet = workbook.active
        all_rows = list(sheet.iter_rows(values_only=True))
    finally:
        workbook.close()

    if not all_rows:
        return []

    header_idx = 0
    headers: list[str] = []
    for idx, row in enumerate(all_rows):
        if row is None:
            continue
        cells = list(row)
        non_empty = [cell for cell in cells if cell not in (None, "")]
        if len(non_empty) >= 2:
            header_idx = idx
            headers = [_normalize_header(value, i) for i, value in enumerate(cells)]
            break

    if not headers:
        headers = [_normalize_header(value, i) for i, value in enumerate(all_rows[0])]
        header_idx = 0

    out: list[dict[str, Any]] = []
    for row in all_rows[header_idx + 1 :]:
        if row is None:
            continue
        values = list(row)
        if all(value in (None, "") for value in values):
            continue
        if len(values) < len(headers):
            values.extend([None] * (len(headers) - len(values)))
        record = {headers[i]: values[i] for i in range(len(headers))}
        out.append(record)
    return out


def read_tabular_rows(path: str | Path) -> list[dict[str, Any]]:
    file_path = Path(path)
    suffix = file_path.suffix.lower()
    if suffix == ".json":
        return _read_json_rows(file_path)
    if suffix == ".csv":
        return _read_csv_rows(file_path, delimiter=",")
    if suffix == ".tsv":
        return _read_csv_rows(file_path, delimiter="\t")
    if suffix in {".xlsx", ".xlsm"}:
        return _read_xlsx_rows(file_path)
    raise RuntimeError(f"unsupported file extension: {suffix}")
